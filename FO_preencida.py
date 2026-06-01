import streamlit as st
import pandas as pd
import io
import tempfile
import os

from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

# ============================================================
# CONFIGURAÇÃO DA PÁGINA
# ============================================================

st.set_page_config(
    page_title="FO SENAI Studio",
    page_icon="📋",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# ============================================================
# CSS PERSONALIZADO
# ============================================================

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@300;400;500;600;700&family=Space+Mono:wght@400;700&display=swap');

:root {
    --azul:     #1a56db;
    --azul-esc: #1239a6;
    --cinza:    #f0f2f6;
    --borda:    #d1d5db;
    --texto:    #1a1a2e;
    --verde:    #0d9e6e;
    --laranja:  #f59e0b;
    --radius:   12px;
}

html, body, [class*="css"] {
    font-family: 'DM Sans', sans-serif !important;
    color: var(--texto);
}

/* ---- cabeçalho ---- */
.header-box {
    background: linear-gradient(135deg, #1a56db 0%, #1239a6 60%, #0d0d2b 100%);
    border-radius: var(--radius);
    padding: 32px 40px;
    margin-bottom: 28px;
    display: flex;
    align-items: center;
    gap: 20px;
    box-shadow: 0 8px 32px rgba(26,86,219,.25);
}
.header-box .icon { font-size: 2.8rem; }
.header-box h1 {
    font-family: 'Space Mono', monospace;
    font-size: 2rem;
    font-weight: 700;
    color: #fff;
    margin: 0; line-height: 1.1;
}
.header-box p {
    color: rgba(255,255,255,.7);
    font-size: .95rem;
    margin: 4px 0 0;
}

/* ---- cards ---- */
.card {
    background: #fff;
    border: 1px solid var(--borda);
    border-radius: var(--radius);
    padding: 24px;
    margin-bottom: 18px;
    box-shadow: 0 2px 8px rgba(0,0,0,.04);
}
.card-title {
    font-family: 'Space Mono', monospace;
    font-size: .8rem;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: .12em;
    color: var(--azul);
    border-bottom: 2px solid var(--cinza);
    padding-bottom: 10px;
    margin-bottom: 16px;
}

/* ---- zona de drop ---- */
.drop-zone {
    border: 2px dashed var(--borda);
    border-radius: var(--radius);
    padding: 32px 20px;
    text-align: center;
    background: var(--cinza);
    transition: border-color .2s, background .2s;
    cursor: pointer;
}
.drop-zone:hover {
    border-color: var(--azul);
    background: #e8f0fd;
}
.drop-zone .dz-icon { font-size: 2.2rem; display: block; margin-bottom: 8px; }
.drop-zone .dz-text { font-size: .95rem; color: #555; }
.drop-zone .dz-hint  { font-size: .8rem; color: #999; margin-top: 4px; }

/* ---- badge de status ---- */
.badge {
    display: inline-flex;
    align-items: center;
    gap: 6px;
    font-size: .82rem;
    font-weight: 600;
    padding: 4px 12px;
    border-radius: 20px;
    margin-top: 8px;
}
.badge-ok  { background: #d1fae5; color: #065f46; }
.badge-err { background: #fee2e2; color: #991b1b; }

/* ---- botão principal ---- */
div[data-testid="stButton"] > button[kind="primary"] {
    background: linear-gradient(135deg, var(--azul), var(--azul-esc));
    border: none;
    border-radius: 8px;
    font-weight: 700;
    letter-spacing: .03em;
    padding: 12px 32px;
    font-size: 1rem;
    box-shadow: 0 4px 14px rgba(26,86,219,.35);
    transition: transform .15s, box-shadow .15s;
}
div[data-testid="stButton"] > button[kind="primary"]:hover {
    transform: translateY(-1px);
    box-shadow: 0 6px 20px rgba(26,86,219,.45);
}

/* ---- inputs ---- */
div[data-testid="stTextInput"] input,
div[data-testid="stTextArea"] textarea {
    border-radius: 8px !important;
    border-color: var(--borda) !important;
    font-family: 'DM Sans', sans-serif !important;
}
div[data-testid="stTextInput"] input:focus,
div[data-testid="stTextArea"] textarea:focus {
    border-color: var(--azul) !important;
    box-shadow: 0 0 0 3px rgba(26,86,219,.12) !important;
}

/* ---- radio ---- */
div[data-testid="stRadio"] label { font-weight: 500; }

/* ---- log ---- */
.log-area {
    background: #0d0d1a;
    border-radius: 8px;
    padding: 16px;
    font-family: 'Space Mono', monospace;
    font-size: .78rem;
    color: #7ee8a2;
    min-height: 120px;
    max-height: 220px;
    overflow-y: auto;
    line-height: 1.7;
    white-space: pre-wrap;
}

/* ---- file uploader override ---- */
div[data-testid="stFileUploaderDropzone"] {
    border-radius: var(--radius) !important;
    border: 2px dashed var(--borda) !important;
    background: var(--cinza) !important;
    padding: 20px !important;
    transition: border-color .2s !important;
}
div[data-testid="stFileUploaderDropzone"]:hover {
    border-color: var(--azul) !important;
    background: #e8f0fd !important;
}
div[data-testid="stFileUploaderDropzone"] svg { color: var(--azul) !important; }

/* ---- hide streamlit padrão ---- */
#MainMenu, footer { visibility: hidden; }
.block-container { padding-top: 1.5rem !important; }
</style>
""", unsafe_allow_html=True)

# ============================================================
# FUNÇÕES DO NEGÓCIO (igual ao original)
# ============================================================

def escrever_celula(celula, texto, alinhamento="left"):
    celula.text = ""
    celula.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragrafo = celula.paragraphs[0]
    paragrafo.alignment = (
        WD_PARAGRAPH_ALIGNMENT.CENTER
        if alinhamento == "center"
        else WD_PARAGRAPH_ALIGNMENT.LEFT
    )
    run = paragrafo.add_run(str(texto))
    run.font.name = "Arial"
    run.font.size = Pt(10)


def extrair_dados_word(arquivo_bytes):
    """Lê o .docx a partir de bytes e devolve o dicionário de dados."""
    dados = {
        "curso": "", "uc": "", "docente": "", "carga": "",
        "funcao": "", "subfuncao": "", "objetivo": "", "estrategia": ""
    }
    try:
        doc = Document(io.BytesIO(arquivo_bytes))

        tabela0 = doc.tables[0]
        dados["curso"]    = tabela0.cell(0, 1).text.strip()
        dados["uc"]       = tabela0.cell(1, 1).text.strip()
        dados["docente"]  = tabela0.cell(2, 1).text.strip()
        dados["carga"]    = tabela0.cell(2, 3).text.strip()

        tabela1 = doc.tables[1]
        dados["funcao"]    = tabela1.cell(1, 1).text.strip()
        dados["subfuncao"] = tabela1.cell(2, 1).text.strip()
        dados["objetivo"]  = tabela1.cell(3, 1).text.strip()

        tabela2 = doc.tables[2]
        if   "(X)" in tabela2.cell(1, 0).text: dados["estrategia"] = "Estudo de Caso"
        elif "(X)" in tabela2.cell(1, 1).text: dados["estrategia"] = "Projeto"
        elif "(X)" in tabela2.cell(1, 2).text: dados["estrategia"] = "Situação-Problema"
        elif "(X)" in tabela2.cell(1, 3).text: dados["estrategia"] = "Pesquisa Aplicada"

    except Exception as e:
        st.error(f"Erro ao ler Word: {e}")

    return dados


def extrair_descricao_excel(arquivo_bytes):
    """
    Le a planilha 'Descricao' usando openpyxl diretamente.

    Estrutura real do Excel (base 1):
      Linha 1 -> titulo da secao
      Linha 2 -> cabecalho: 'Etapa da Estrategia' | 'Detalhamento...'
      Linha 3 -> 'CONTEXTUALIZACAO:' | texto (coluna B)
      Linha 4 -> 'DESAFIO:'          | texto (coluna B)
      Linha 5 -> 'RESULTADOS ESPERADOS:' | texto (coluna B)

    Os valores em B3, B4 e B5 comecam com '\t\n' (tab + newline)
    inseridos pelo Excel como prefixo de formatacao -- sao removidos aqui.
    """
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(arquivo_bytes), data_only=True)
        ws = wb["Descricao"]

        def _limpar(valor):
            if valor is None:
                return ""
            return str(valor).lstrip("\t\n\r ").strip()

        contextualizacao = _limpar(ws.cell(row=3, column=2).value)
        desafio          = _limpar(ws.cell(row=4, column=2).value)
        resultados       = _limpar(ws.cell(row=5, column=2).value)
        wb.close()

        return {
            "contextualizacao": contextualizacao,
            "desafio":          desafio,
            "resultados":       resultados,
        }

    except Exception as e:
        st.error(f"Erro ao ler descricao: {e}")
        return {"contextualizacao": "", "desafio": "", "resultados": ""}

def ler_plano_excel(arquivo_bytes):
    try:
        return pd.read_excel(
            io.BytesIO(arquivo_bytes),
            sheet_name="Plano Semanal",
            skiprows=2
        )
    except Exception as e:
        st.error(f"Erro ao ler plano: {e}")
        return None


def inspecionar_word(arquivo_bytes):
    doc = Document(io.BytesIO(arquivo_bytes))
    linhas = []
    for t, tabela in enumerate(doc.tables):
        linhas.append("\n" + "=" * 60)
        linhas.append(f"TABELA {t}")
        linhas.append("=" * 60)
        for r, linha in enumerate(tabela.rows):
            for c, cel in enumerate(linha.cells):
                linhas.append(f"[{r},{c}] {cel.text.strip()}")
    return "\n".join(linhas)


def gerar_fo(word_bytes, excel_bytes, form):
    """Gera o FO preenchido e retorna os bytes do arquivo resultante."""
    doc = Document(io.BytesIO(word_bytes))

    # Tabela 0
    tabela0 = doc.tables[0]
    escrever_celula(tabela0.cell(2, 1), form["docente"].upper())

    # Tabela 1
    tabela1 = doc.tables[1]
    escrever_celula(tabela1.cell(1, 1), form["funcao"].upper())
    escrever_celula(tabela1.cell(2, 1), form["subfuncao"].upper())
    escrever_celula(tabela1.cell(3, 1), form["objetivo"])

    # Tabela 2 — estratégia
    tabela2 = doc.tables[2]
    textos = [
        "( ) Estudo de caso",
        "( ) Projeto (elaboração ou execução)",
        "( ) Situação-Problema",
        "( ) Pesquisa Aplicada",
    ]
    for i in range(4):
        tabela2.cell(1, i).text = textos[i]

    mapa = {
        "Estudo de Caso": 0,
        "Projeto": 1,
        "Situação-Problema": 2,
        "Pesquisa Aplicada": 3,
    }
    idx = mapa.get(form["estrategia"], 1)
    tabela2.cell(1, idx).text = tabela2.cell(1, idx).text.replace("( )", "(X)")

    # Tabela 3 - Descricao da Estrategia de Aprendizagem Desafiadora
    #
    # Estrutura original: 7 linhas. Cada secao (Contextualizacao, Desafio,
    # Resultados) ocupa 2 linhas via vMerge vertical:
    #   - Linha master (restart): trHeight=274, colB bottom=nil (borda aberta)
    #   - Linha continuation:     trHeight=410, colB top=nil    (borda aberta)
    #
    # Ao remover as linhas de continuacao, as bordas ficam abertas visualmente.
    # Correcoes necessarias em cada linha master apos remocao:
    #   1. Remover vMerge da coluna A
    #   2. Ajustar trHeight para 274+410=684 (soma das duas linhas originais)
    #   3. Coluna A: setar bottom=single (fechar borda inferior)
    #   4. Coluna B: setar bottom=single (estava nil por causa do merge)
    tabela3 = doc.tables[3]
    desc = extrair_descricao_excel(excel_bytes)

    W  = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    tbl_xml = tabela3._tbl
    todas_tr = tbl_xml.findall(f"{{{W}}}tr")

    # --- Helper: criar/substituir elemento de borda ---
    def _set_border(tcPr, lado, val="single", sz="4", color="000000"):
        """Garante que a borda `lado` da celula esteja com os valores corretos."""
        from lxml import etree
        tcBorders = tcPr.find(f"{{{W}}}tcBorders")
        if tcBorders is None:
            tcBorders = etree.SubElement(tcPr, f"{{{W}}}tcBorders")
        borda = tcBorders.find(f"{{{W}}}"+lado)
        if borda is None:
            borda = etree.SubElement(tcBorders, f"{{{W}}}"+lado)
        borda.set(f"{{{W}}}val",   val)
        borda.set(f"{{{W}}}sz",    sz)
        borda.set(f"{{{W}}}space", "0")
        borda.set(f"{{{W}}}color", color)

    # --- Corrigir cada linha master (indices 1, 3, 5) antes de remover ---
    # As linhas de continuacao estao nos indices 2, 4, 6
    for master_idx in (1, 3, 5):
        tr_master = todas_tr[master_idx]
        cont_idx  = master_idx + 1
        tr_cont   = todas_tr[cont_idx]

        # 1. Somar alturas (trHeight master + continuation)
        trPr_m = tr_master.find(f"{{{W}}}trPr")
        trPr_c = tr_cont.find(f"{{{W}}}trPr")
        h_m = int(trPr_m.find(f"{{{W}}}trHeight").get(f"{{{W}}}val", "274"))
        h_c = int(trPr_c.find(f"{{{W}}}trHeight").get(f"{{{W}}}val", "410"))
        trPr_m.find(f"{{{W}}}trHeight").set(f"{{{W}}}val", str(h_m + h_c))

        tcs = tr_master.findall(f"{{{W}}}tc")

        # 2. Coluna A: remover vMerge e fechar borda inferior
        tcPr_a = tcs[0].find(f"{{{W}}}tcPr")
        vmerge = tcPr_a.find(f"{{{W}}}vMerge")
        if vmerge is not None:
            tcPr_a.remove(vmerge)
        _set_border(tcPr_a, "bottom")

        # 3. Coluna B: fechar borda inferior (estava nil) e garantir top
        tcPr_b = tcs[1].find(f"{{{W}}}tcPr")
        _set_border(tcPr_b, "bottom")
        _set_border(tcPr_b, "top")

    # --- Remover linhas de continuacao (de tras pra frente) ---
    for idx in (6, 4, 2):
        tbl_xml.remove(todas_tr[idx])

    # --- Escrever conteudo nas linhas corretas (1, 2, 3 apos remocao) ---
    escrever_celula(tabela3.cell(1, 1), desc["contextualizacao"])
    escrever_celula(tabela3.cell(2, 1), desc["desafio"])
    escrever_celula(tabela3.cell(3, 1), desc["resultados"])

    # Tabela 4 — plano semanal
    df = ler_plano_excel(excel_bytes)
    if df is not None:
        tabela4 = doc.tables[4]
        for i, (_, linha) in enumerate(df.iterrows(), start=1):
            if i >= len(tabela4.rows):
                break
            escrever_celula(tabela4.rows[i].cells[0],
                            f"{linha['Aula']}\n{linha['Bimestre']}", "center")
            escrever_celula(tabela4.rows[i].cells[1], linha["CH"], "center")
            escrever_celula(tabela4.rows[i].cells[2],
                            linha["Fundamentos Técnicos / Capacidades"])
            escrever_celula(tabela4.rows[i].cells[3], linha["Conhecimentos"])
            escrever_celula(tabela4.rows[i].cells[4], linha["Estratégias de Ensino"])
            escrever_celula(tabela4.rows[i].cells[5], linha["Critérios de Avaliação"])
            escrever_celula(tabela4.rows[i].cells[6],
                            linha["Instrumentos de Avaliação"])
            escrever_celula(tabela4.rows[i].cells[7],
                            linha["Recursos Didáticos, Ambientes e Acessibilidade"])

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

# ============================================================
# SESSION STATE
# ============================================================

for key in ("word_bytes", "excel_bytes", "log_lines", "dados_word"):
    if key not in st.session_state:
        st.session_state[key] = None if key != "log_lines" else []

def adicionar_log(msg):
    st.session_state.log_lines.append(f"› {msg}")

# ============================================================
# INTERFACE
# ============================================================

# Cabeçalho
st.markdown("""
<div class="header-box">
  <span class="icon">📋</span>
  <div>
    <h1>FO SENAI Studio</h1>
    <p>Gerador de Fichas de Operação · versão Web</p>
  </div>
</div>
""", unsafe_allow_html=True)

# -------- Layout: duas colunas principais --------
col_esq, col_dir = st.columns([1, 1.6], gap="large")

# ============================================================
# COLUNA ESQUERDA — upload de arquivos
# ============================================================
with col_esq:

    # ---- Word ----
    st.markdown('<div class="card"><div class="card-title">📄 Arquivo Word</div>', unsafe_allow_html=True)
    word_up = st.file_uploader(
        "Arraste o .docx aqui ou clique para selecionar",
        type=["docx"],
        key="word_uploader",
        label_visibility="visible"
    )
    if word_up:
        st.session_state.word_bytes = word_up.read()
        st.session_state.dados_word = extrair_dados_word(st.session_state.word_bytes)
        adicionar_log(f"Word carregado: {word_up.name}")
        st.markdown('<span class="badge badge-ok">✓ Arquivo carregado</span>', unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)

    # ---- Excel ----
    st.markdown('<div class="card"><div class="card-title">📊 Arquivo Excel</div>', unsafe_allow_html=True)
    excel_up = st.file_uploader(
        "Arraste o .xlsx aqui ou clique para selecionar",
        type=["xlsx"],
        key="excel_uploader",
        label_visibility="visible"
    )
    if excel_up:
        st.session_state.excel_bytes = excel_up.read()
        plano = ler_plano_excel(st.session_state.excel_bytes)
        qtd = len(plano) if plano is not None else 0
        adicionar_log(f"Excel carregado: {excel_up.name}")
        adicionar_log(f"{qtd} aulas encontradas no plano semanal.")
        st.markdown(f'<span class="badge badge-ok">✓ {qtd} aulas encontradas</span>', unsafe_allow_html=True)

        # Diagnóstico: pré-visualiza os 3 campos para confirmar leitura correta
        desc_preview = extrair_descricao_excel(st.session_state.excel_bytes)
        with st.expander("🔎 Pré-visualizar campos da planilha Descricao"):
            for campo, label in [
                ("contextualizacao", "Contextualização"),
                ("desafio",          "Desafio"),
                ("resultados",       "Resultados Esperados"),
            ]:
                valor = desc_preview[campo]
                chars = len(valor)
                cor = "#0d9e6e" if chars > 20 else "#f59e0b"
                st.markdown(
                    f'<div style="margin-bottom:10px">'
                    f'<span style="font-weight:700;color:{cor};">{label}</span> '
                    f'<span style="font-size:.78rem;color:#999;">({chars} caracteres)</span>'
                    f'<div style="background:#f0f2f6;border-radius:6px;padding:8px 12px;'
                    f'font-size:.82rem;margin-top:4px;white-space:pre-wrap;">'
                    f'{valor[:300]}{"..." if chars > 300 else ""}'
                    f'</div></div>',
                    unsafe_allow_html=True
                )
            if not any(desc_preview.values()):
                st.warning("Nenhum campo foi lido. Verifique se a planilha se chama exatamente **Descricao** e se os dados estão nas linhas 3, 4 e 5 da coluna B.")
            else:
                adicionar_log(
                    f"Descricao lida: Contextualizacao={len(desc_preview['contextualizacao'])}c / "
                    f"Desafio={len(desc_preview['desafio'])}c / "
                    f"Resultados={len(desc_preview['resultados'])}c"
                )
    st.markdown('</div>', unsafe_allow_html=True)

    # ---- Log ----
    st.markdown('<div class="card"><div class="card-title">⚡ Log</div>', unsafe_allow_html=True)
    log_texto = "\n".join(st.session_state.log_lines) if st.session_state.log_lines else "Aguardando arquivos..."
    st.markdown(f'<div class="log-area">{log_texto}</div>', unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)

# ============================================================
# COLUNA DIREITA — formulário
# ============================================================
with col_dir:

    dados = st.session_state.dados_word or {}

    st.markdown('<div class="card"><div class="card-title">🎓 Identificação</div>', unsafe_allow_html=True)

    c1, c2 = st.columns(2)
    with c1:
        curso    = st.text_input("Curso",    value=dados.get("curso",    ""))
        docente  = st.text_input("Docente",  value=dados.get("docente",  ""))
        funcao   = st.text_input("Função",   value=dados.get("funcao",   ""))
    with c2:
        uc       = st.text_input("Unidade Curricular", value=dados.get("uc",       ""))
        carga    = st.text_input("Carga Horária",       value=dados.get("carga",    ""), disabled=True)
        subfuncao = st.text_input("Subfunção",          value=dados.get("subfuncao",""))

    st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('<div class="card"><div class="card-title">🎯 Objetivo & Estratégia</div>', unsafe_allow_html=True)

    objetivo = st.text_area(
        "Objetivo Geral",
        value=dados.get("objetivo", ""),
        height=130
    )

    estrategia = st.radio(
        "Estratégia de Aprendizagem",
        options=["Estudo de Caso", "Projeto", "Situação-Problema", "Pesquisa Aplicada"],
        index=["Estudo de Caso", "Projeto", "Situação-Problema", "Pesquisa Aplicada"].index(
            dados.get("estrategia", "Projeto")
        ) if dados.get("estrategia") else 1,
        horizontal=True
    )

    st.markdown('</div>', unsafe_allow_html=True)

    # ---- Ações ----
    st.markdown('<div class="card"><div class="card-title">⚙️ Ações</div>', unsafe_allow_html=True)

    col_btn1, col_btn2 = st.columns(2)

    with col_btn1:
        if st.button("🔍 Inspecionar Word", use_container_width=True):
            if st.session_state.word_bytes:
                resultado = inspecionar_word(st.session_state.word_bytes)
                with st.expander("Estrutura das tabelas do Word", expanded=True):
                    st.code(resultado, language=None)
            else:
                st.warning("Carregue o arquivo Word primeiro.")

    with col_btn2:
        gerar_btn = st.button("🚀 Gerar FO", type="primary", use_container_width=True)

    if gerar_btn:
        if not st.session_state.word_bytes:
            st.error("⚠️ Selecione o arquivo Word.")
        elif not st.session_state.excel_bytes:
            st.error("⚠️ Selecione o arquivo Excel.")
        else:
            with st.spinner("Gerando Ficha de Operação..."):
                form = {
                    "curso":     curso,
                    "uc":        uc,
                    "docente":   docente,
                    "funcao":    funcao,
                    "subfuncao": subfuncao,
                    "objetivo":  objetivo,
                    "estrategia": estrategia,
                }
                try:
                    fo_bytes = gerar_fo(
                        st.session_state.word_bytes,
                        st.session_state.excel_bytes,
                        form
                    )
                    adicionar_log("FO gerada com sucesso!")
                    st.success("✅ Ficha de Operação gerada! Clique abaixo para baixar.")
                    st.download_button(
                        label="⬇️ Baixar FO Preenchida (.docx)",
                        data=fo_bytes,
                        file_name="FO_Preenchida.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True
                    )
                except Exception as e:
                    st.error(f"Erro ao gerar FO: {e}")
                    adicionar_log(f"ERRO: {e}")

    st.markdown('</div>', unsafe_allow_html=True)
